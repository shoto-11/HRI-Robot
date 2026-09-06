"""Export risk appearance chart to PowerPoint (.pptx) and PDF."""
from __future__ import annotations

import colorsys
import os
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas as pdf_canvas

OUT = Path(__file__).resolve().parent
SAMPLES = [0.0, 0.25, 0.5, 0.75, 1.0]
LOOKS = ["青緑・うすい", "緑・ややうすい", "黄緑・中くらい", "橙黄・やや濃い", "赤・はっきり"]


def map_r(r: float):
    h = (1 - r) * 180 / 360
    s = 0.4 + 0.6 * r
    v = 0.5 + 0.3 * r
    a = 0.35 + 0.65 * r
    rgb = colorsys.hsv_to_rgb(h, s, v)
    return h * 360, s, v, a, tuple(int(c * 255) for c in rgb)


def font(size: int):
    for p in [
        r"C:\Windows\Fonts\meiryo.ttc",
        r"C:\Windows\Fonts\YuGothM.ttc",
        r"C:\Windows\Fonts\msgothic.ttc",
        r"C:\Windows\Fonts\arial.ttf",
    ]:
        if os.path.exists(p):
            try:
                return ImageFont.truetype(p, size, index=0)
            except Exception:
                continue
    return ImageFont.load_default()


def paint_alpha(img: Image.Image, box, rgb, a, tile=14):
    x0, y0, x1, y1 = box
    R, G, B = rgb
    for y in range(y0, y1):
        for x in range(x0, x1):
            bg = 220 if ((x // tile) + (y // tile)) % 2 == 0 else 255
            img.putpixel(
                (x, y),
                (
                    int(R * a + bg * (1 - a)),
                    int(G * a + bg * (1 - a)),
                    int(B * a + bg * (1 - a)),
                ),
            )


def ensure_slide_png() -> Path:
    path = OUT / "risk_appearance_slide_16x9.png"
    if path.exists():
        return path

    W, H = 1920, 1080
    img = Image.new("RGB", (W, H), (255, 255, 255))
    draw = ImageDraw.Draw(img)
    f_title, f_sub, f_lab, f_num = font(44), font(22), font(26), font(20)
    draw.text((80, 48), "危険度スコア R ごとの経路表示（色・不透明度）", fill=(20, 20, 20), font=f_title)
    draw.text(
        (80, 110),
        "α = 0.35 + 0.65R    H = (1−R)×180°    S = 0.4+0.6R    V = 0.5+0.3R",
        fill=(90, 90, 90),
        font=f_sub,
    )
    draw.text(
        (80, 150),
        "上段：色のみ（不透明）　／　下段：不透明度込み（市松模様の上）",
        fill=(90, 90, 90),
        font=f_sub,
    )
    margin, gap, n = 80, 28, 5
    cell_w = (W - 2 * margin - gap * (n - 1)) // n
    y_top, sw_h = 220, 120
    for i, r in enumerate(SAMPLES):
        Hdeg, S, V, a, rgb = map_r(r)
        x = margin + i * (cell_w + gap)
        draw.rectangle([x, y_top, x + cell_w, y_top + sw_h], fill=rgb, outline=(30, 30, 30), width=2)
        box = (x, y_top + sw_h + 16, x + cell_w, y_top + 2 * sw_h + 16)
        paint_alpha(img, box, rgb, a)
        draw.rectangle(list(box), outline=(30, 30, 30), width=2)
        ty = box[3] + 24
        for j, t in enumerate(
            [
                f"R = {r:g}",
                LOOKS[i],
                f"H = {Hdeg:.0f}°",
                f"α = {a:.3f}",
                f"S = {S:.2f}  V = {V:.2f}",
                f"#{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}",
            ]
        ):
            draw.text(
                (x + 8, ty + j * 32),
                t,
                fill=(10, 10, 10) if j == 0 else (50, 50, 50),
                font=f_lab if j <= 1 else f_num,
            )
    img.save(path, "PNG")
    return path


def export_pptx(slide_png: Path) -> Path:
    prs = Presentation()
    # 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Slide 1: overview image full-bleed
    blank = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank)
    s1.shapes.add_picture(str(slide_png), Inches(0), Inches(0), width=prs.slide_width)

    # Slide 2: table of values + color boxes
    s2 = prs.slides.add_slide(blank)
    title = s2.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.6))
    tf = title.text_frame
    p = tf.paragraphs[0]
    p.text = "危険度 R ごとの色相・不透明度（数値）"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(20, 20, 20)

    rows, cols = 6, 7
    table_shape = s2.shapes.add_table(rows, cols, Inches(0.5), Inches(1.1), Inches(12.3), Inches(3.2))
    table = table_shape.table
    headers = ["R", "見た目", "色相 H", "彩度 S", "明度 V", "不透明度 α", "RGB"]
    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        for para in cell.text_frame.paragraphs:
            para.font.bold = True
            para.font.size = Pt(14)

    for i, r in enumerate(SAMPLES):
        Hdeg, S, V, a, rgb = map_r(r)
        vals = [
            f"{r:g}",
            LOOKS[i],
            f"{Hdeg:.0f}°",
            f"{S:.2f}",
            f"{V:.2f}",
            f"{a:.3f}",
            f"#{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}",
        ]
        for c, v in enumerate(vals):
            cell = table.cell(i + 1, c)
            cell.text = v
            for para in cell.text_frame.paragraphs:
                para.font.size = Pt(13)

    # color chips under table
    chip_y = Inches(4.6)
    chip_w = Inches(2.2)
    gap = Inches(0.2)
    x0 = Inches(0.5)
    for i, r in enumerate(SAMPLES):
        _, _, _, a, rgb = map_r(r)
        left = x0 + i * (chip_w + gap)
        shape = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, chip_y, chip_w, Inches(1.0))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(*rgb)
        shape.line.color.rgb = RGBColor(40, 40, 40)
        label = s2.shapes.add_textbox(left, chip_y + Inches(1.1), chip_w, Inches(0.8))
        lf = label.text_frame
        lf.paragraphs[0].text = f"R={r:g}"
        lf.paragraphs[0].font.size = Pt(14)
        lf.paragraphs[0].font.bold = True
        lf.paragraphs[0].alignment = PP_ALIGN.CENTER
        p2 = lf.add_paragraph()
        p2.text = f"α={a:.3f}"
        p2.font.size = Pt(12)
        p2.alignment = PP_ALIGN.CENTER

    # Slides 3–7: one per R
    for i, r in enumerate(SAMPLES):
        Hdeg, S, V, a, rgb = map_r(r)
        s = prs.slides.add_slide(blank)
        t = s.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(0.7))
        tp = t.text_frame.paragraphs[0]
        tp.text = f"R = {r:g}　（{LOOKS[i]}）"
        tp.font.size = Pt(32)
        tp.font.bold = True

        # big color rect
        rect = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(1.5), Inches(10), Inches(2.5))
        rect.fill.solid()
        rect.fill.fore_color.rgb = RGBColor(*rgb)
        rect.line.color.rgb = RGBColor(30, 30, 30)

        body = s.shapes.add_textbox(Inches(1.5), Inches(4.3), Inches(10), Inches(2.5))
        bf = body.text_frame
        lines = [
            f"色相 H = {Hdeg:.0f}°",
            f"彩度 S = {S:.2f}",
            f"明度 V = {V:.2f}",
            f"不透明度 α = {a:.3f}  （α = 0.35 + 0.65R）",
            f"RGB = ({rgb[0]}, {rgb[1]}, {rgb[2]})  #{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}",
        ]
        bf.paragraphs[0].text = lines[0]
        bf.paragraphs[0].font.size = Pt(20)
        for line in lines[1:]:
            para = bf.add_paragraph()
            para.text = line
            para.font.size = Pt(20)

    out = OUT / "risk_appearance.pptx"
    prs.save(out)
    return out


def register_jp_font() -> str:
    for name, path in [
        ("Meiryo", r"C:\Windows\Fonts\meiryo.ttc"),
        ("YuGothic", r"C:\Windows\Fonts\YuGothM.ttc"),
        ("MSGothic", r"C:\Windows\Fonts\msgothic.ttc"),
    ]:
        if os.path.exists(path):
            try:
                pdfmetrics.registerFont(TTFont(name, path, subfontIndex=0))
                return name
            except Exception:
                continue
    return "Helvetica"


def export_pdf(slide_png: Path) -> Path:
    out = OUT / "risk_appearance.pdf"
    # landscape 16:9-ish page
    page = (1920, 1080)
    c = pdf_canvas.Canvas(str(out), pagesize=page)
    font_name = register_jp_font()

    # Page 1: full image
    c.drawImage(ImageReader(str(slide_png)), 0, 0, width=page[0], height=page[1])
    c.showPage()

    # Page 2: table
    c.setFont(font_name, 28)
    c.drawString(80, page[1] - 80, "危険度 R ごとの色相・不透明度（数値）")
    c.setFont(font_name, 16)
    headers = ["R", "見た目", "H", "S", "V", "α", "RGB"]
    xs = [80, 180, 420, 560, 680, 800, 980]
    y = page[1] - 160
    for x, h in zip(xs, headers):
        c.drawString(x, y, h)
    y -= 40
    for i, r in enumerate(SAMPLES):
        Hdeg, S, V, a, rgb = map_r(r)
        vals = [
            f"{r:g}",
            LOOKS[i],
            f"{Hdeg:.0f}°",
            f"{S:.2f}",
            f"{V:.2f}",
            f"{a:.3f}",
            f"#{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}",
        ]
        for x, v in zip(xs, vals):
            c.drawString(x, y, v)
        # color chip
        c.setFillColorRGB(rgb[0] / 255, rgb[1] / 255, rgb[2] / 255)
        c.rect(1400, y - 8, 120, 28, fill=1, stroke=1)
        c.setFillColorRGB(0, 0, 0)
        y -= 48

    c.setFont(font_name, 14)
    c.drawString(80, 80, "HRI-Robot / RiskToVisualMapper")
    c.showPage()

    # Pages per R
    for i, r in enumerate(SAMPLES):
        Hdeg, S, V, a, rgb = map_r(r)
        c.setFont(font_name, 32)
        c.drawString(80, page[1] - 100, f"R = {r:g}　（{LOOKS[i]}）")
        c.setFillColorRGB(rgb[0] / 255, rgb[1] / 255, rgb[2] / 255)
        c.rect(200, 420, 1520, 280, fill=1, stroke=1)
        c.setFillColorRGB(0, 0, 0)
        c.setFont(font_name, 22)
        lines = [
            f"色相 H = {Hdeg:.0f}°",
            f"彩度 S = {S:.2f}　　明度 V = {V:.2f}",
            f"不透明度 α = {a:.3f}  （α = 0.35 + 0.65R）",
            f"RGB = ({rgb[0]}, {rgb[1]}, {rgb[2]})  #{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}",
        ]
        yy = 340
        for line in lines:
            c.drawString(200, yy, line)
            yy -= 40
        c.showPage()

    c.save()
    return out


def main():
    slide_png = ensure_slide_png()
    pptx = export_pptx(slide_png)
    pdf = export_pdf(slide_png)
    print("pptx:", pptx)
    print("pdf:", pdf)


if __name__ == "__main__":
    main()
