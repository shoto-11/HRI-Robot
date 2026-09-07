"""One-slide summary of risk score R and visual mapping formulas."""
from __future__ import annotations

import os
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches

OUT = Path(__file__).resolve().parent
W, H = 1920, 1080


def jp_font(size: int):
    for p in [
        r"C:\Windows\Fonts\meiryo.ttc",
        r"C:\Windows\Fonts\YuGothM.ttc",
        r"C:\Windows\Fonts\msgothic.ttc",
    ]:
        if os.path.exists(p):
            try:
                return ImageFont.truetype(p, size, index=0)
            except Exception:
                continue
    return ImageFont.load_default()


def rounded_rect(draw, box, fill, outline=(200, 200, 200), radius=16, width=2):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def draw_formula_slide_png() -> Path:
    img = Image.new("RGB", (W, H), (250, 250, 252))
    d = ImageDraw.Draw(img)
    f_title = jp_font(40)
    f_h = jp_font(24)
    f_body = jp_font(20)
    f_eq = jp_font(22)
    f_small = jp_font(16)
    f_tiny = jp_font(14)

    d.rectangle([0, 0, W, 88], fill=(32, 56, 88))
    d.text((48, 22), "危険度スコア R の計算と視覚マッピング（Proposed）", fill=(255, 255, 255), font=f_title)

    def card(xyxy, title, lines, title_bg=(45, 90, 140)):
        x0, y0, x1, y1 = xyxy
        rounded_rect(d, xyxy, (255, 255, 255), (210, 215, 225), 14)
        d.rounded_rectangle([x0, y0, x1, y0 + 44], radius=14, fill=title_bg)
        d.rectangle([x0, y0 + 22, x1, y0 + 44], fill=title_bg)
        d.text((x0 + 18, y0 + 10), title, fill=(255, 255, 255), font=f_h)
        yy = y0 + 58
        for kind, text in lines:
            f = f_eq if kind == "eq" else (f_small if kind == "small" else f_body)
            color = (25, 25, 30) if kind != "small" else (80, 85, 95)
            d.text((x0 + 18, yy), text, fill=color, font=f)
            yy += 30 if kind != "eq" else 34

    card(
        (40, 110, 940, 250),
        "① 表示可否（距離のみ）",
        [
            ("eq", "visible = ( d ≤ Dmax )"),
            ("body", "非表示なら R = 0　／　Dmax = 13.6 m　／　Tmax は表示に使わない"),
            ("small", "Dmax = (vAGV,max + vped) × Tmax = (2.0 + 1.4) × 4"),
        ],
        (40, 110, 160),
    )
    card(
        (40, 268, 940, 500),
        "② 三要因（PTTC / Crossing TTC・CTTC / 近接）",
        [
            ("eq", "Tp = d / vclose     vclose = max(0, r̂ · v)（近づかない→∞）"),
            ("eq", "Tc = s / max(vAGV, 0.1)     s: 経路→視線太線までの道のり（CTTC）"),
            ("eq", "Rp,Rc = (1 − T/Tmax)^γ     Rd = (1 − d/Dmax)^γ"),
            ("small", "Tmax = 4.0 s　／　γ = 0.6　／　基準線半幅 w = 0.5 m"),
        ],
        (30, 120, 100),
    )
    card(
        (40, 518, 940, 700),
        "③ 統合スコア R（重み付き和）",
        [
            ("eq", "R = min(1,  wp·Rp + wc·Rc + wd·Rd)"),
            ("body", "wp=0.55（迫り）　wc=0.30（視線横断）　wd=0.15（近接は保険）"),
            ("small", "PTTC は相対速度 vAGV−vped　／　Rfloor なし"),
        ],
        (140, 50, 50),
    )
    card(
        (40, 718, 940, 860),
        "④ 狙い",
        [
            ("body", "遠い正面迫り > 近いが向かない　／　視線横断は Rc で残す"),
            ("small", "実装: VehicleRiskCalculator → RiskToVisualMapper → PathRenderer"),
        ],
        (50, 70, 90),
    )

    card(
        (980, 110, 1880, 430),
        "⑤ 色・不透明度（連続）",
        [
            ("eq", "H(R) = (1 − R) × 180°     青緑(安全) → 赤(危険)"),
            ("eq", "S(R) = 0.4 + 0.6 R"),
            ("eq", "V(R) = 0.5 + 0.3 R"),
            ("eq", "α(R) = 0.35 + 0.65 R      半透明 → 不透明"),
            ("body", "Proposed 条件のみ適用　／　色も不透明度も R に対して連続変化"),
            ("small", "Baseline: 固定青・α=1　／　No-AR: 経路非表示"),
        ],
        (70, 50, 130),
    )

    import colorsys

    bar_x0, bar_y0, bar_x1, bar_y1 = 1000, 460, 1860, 540
    d.text((1000, 448), "R: 0 → 1 の見た目", fill=(60, 60, 70), font=f_small)
    for x in range(bar_x0, bar_x1):
        t = (x - bar_x0) / (bar_x1 - bar_x0)
        h = (1 - t) * 180 / 360
        s = 0.4 + 0.6 * t
        v = 0.5 + 0.3 * t
        a = 0.35 + 0.65 * t
        r, g, b = colorsys.hsv_to_rgb(h, s, v)
        rr = int((r * a + 1 * (1 - a)) * 255)
        gg = int((g * a + 1 * (1 - a)) * 255)
        bb = int((b * a + 1 * (1 - a)) * 255)
        d.line([(x, bar_y0), (x, bar_y1)], fill=(rr, gg, bb))
    d.rectangle([bar_x0, bar_y0, bar_x1, bar_y1], outline=(80, 80, 90), width=2)
    d.text((bar_x0, bar_y1 + 8), "R=0  青緑・うすい", fill=(50, 50, 60), font=f_tiny)
    d.text((bar_x1 - 160, bar_y1 + 8), "R=1  赤・はっきり", fill=(50, 50, 60), font=f_tiny)

    card(
        (980, 590, 1880, 860),
        "主要パラメータ",
        [
            ("body", "Tmax = 4.0 s（正規化のみ）　Dmax = 13.6 m（表示＋Rd）"),
            ("body", "γ = 0.6　wp/wc/wd = 0.55 / 0.30 / 0.15（相対PTTC）"),
            ("body", "vAGV,max = 2.0 m/s　vped = 1.4 m/s　w = 0.5 m"),
            ("small", "表示ゲートは距離のみ　／　スコアは Tp+Tc+d の重み付き和"),
        ],
        (50, 70, 90),
    )

    d.rectangle([0, 1000, W, H], fill=(235, 238, 245))
    d.text(
        (48, 1025),
        "HRI-Robot  Proposed eHMI  ／  危険度計算の1枚まとめ（PTTC + CTTC + 近接）",
        fill=(70, 75, 90),
        font=f_small,
    )

    out = OUT / "risk_formula_slide_16x9.png"
    img.save(out, "PNG")
    return out


def export_pptx(png: Path) -> Path | None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.shapes.add_picture(str(png), Inches(0), Inches(0), width=prs.slide_width)
    out = OUT / "risk_formula_slide.pptx"
    try:
        prs.save(out)
        return out
    except PermissionError:
        alt = OUT / "risk_formula_slide_new.pptx"
        prs.save(alt)
        print("pptx locked; wrote", alt)
        return alt


def main():
    png = draw_formula_slide_png()
    pptx = export_pptx(png)
    print("png:", png)
    print("pptx:", pptx)


if __name__ == "__main__":
    main()
